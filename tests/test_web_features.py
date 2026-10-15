import http.client
import io
import json
import sqlite3
import threading
from datetime import date, datetime, timedelta, timezone
from urllib.parse import quote

import pytest
from openpyxl import Workbook

from collectos import web
from collectos.importer import import_rows
from collectos.store import Store

HEADER = ["Invoice No", "Customer", "Invoice Date", "Due Date", "Amount", "Paid"]
NOW = datetime(2026, 10, 15, 9, 0, tzinfo=timezone.utc)
ASOF = "as_of=2026-10-15"


def make_db(tmp_path):
    path = tmp_path / "t.db"
    rows = [(1, HEADER),
            (2, ["INV-1001", "Acme Traders", "01-06-2026", "01-07-2026", "100000", "0"]),
            (3, ["INV-1002", "Acme Traders", "01-08-2026", "01-09-2026", "40000", "0"]),
            (4, ["INV-2001", "Beta Co", "01-08-2026", "01-09-2026", "20000", "5000"]),
            (5, ["INV-3001", "Gamma Ltd", "01-09-2026", "01-10-2026", "7000", "0"])]
    with Store(path, now=lambda: NOW) as s:
        s.apply_import(import_rows(rows), actor="t")
    return path


@pytest.fixture
def server(tmp_path):
    srv = web.make_server(make_db(tmp_path))
    threading.Thread(target=lambda: srv.serve_forever(poll_interval=0.02), daemon=True).start()
    yield srv
    srv.shutdown(); srv.server_close()


def call(srv, path, method="GET", body=None, headers=None, raw=None, host=None):
    conn = http.client.HTTPConnection("127.0.0.1", srv.server_address[1], timeout=15)
    h = {"Host": host or f"127.0.0.1:{srv.server_address[1]}", **(headers or {})}
    payload = raw if raw is not None else (None if body is None else json.dumps(body).encode())
    conn.putrequest(method, path, skip_host=True)
    for k, v in h.items():
        conn.putheader(k, v)
    if payload is not None:
        conn.putheader("Content-Length", str(len(payload)))
    conn.endheaders(payload)
    resp = conn.getresponse()
    data = resp.read()
    conn.close()
    try:
        parsed = json.loads(data)
    except ValueError:
        parsed = data
    return resp.status, {k.lower(): v for k, v in resp.getheaders()}, parsed


def token(srv):
    return call(srv, "/api/session")[2]["token"]


def post(srv, path, body, **kw):
    return call(srv, path, "POST", body, {"Content-Type": "application/json", web.TOKEN_HEADER: token(srv)}, **kw)


def upload(srv, path, data: bytes, filename="x.csv", ctype="application/octet-stream", tok=None):
    headers = {"Content-Type": ctype, web.TOKEN_HEADER: token(srv) if tok is None else tok, "X-Filename": quote(filename)}
    return call(srv, path, "POST", headers=headers, raw=data)


def overview(srv):
    return call(srv, f"/api/overview?{ASOF}")[2]


def csv_bytes(*rows):
    out = ",".join(HEADER) + "\n"
    return (out + "\n".join(",".join(r) for r in rows) + "\n").encode()


ROW1 = ["INV-1001", "Acme Traders", "01-06-2026", "01-07-2026", "100000", "0"]
ROW2 = ["INV-1002", "Acme Traders", "01-08-2026", "01-09-2026", "40000", "0"]
ROWB = ["INV-2001", "Beta Co", "01-08-2026", "01-09-2026", "20000", "5000"]
ROWG = ["INV-3001", "Gamma Ltd", "01-09-2026", "01-10-2026", "7000", "0"]


# ---- overview additions -----------------------------------------------------------------------------


def test_overview_has_the_new_measures_and_queue_fields(server):
    d = overview(server)
    k = d["kpi"]
    assert k["customers_overdue"] == 3 and k["contacted_7d"] == 0 and k["coverage_7d"] == "0.0000"
    assert k["avg_days_overdue"] is not None and k["promised_7d_count"] == 0 and k["disputed_count"] == 0
    q = d["queue"][0]
    assert q["customer"] == "Acme Traders" and q["stage"] == 4 and q["stage_label"] == "Review credit terms"
    assert "days overdue" in q["stage_why"] and q["contact"]["phone"] == "" and q["promises_kept"] == 0
    assert d["health"]["no_due_date"]["count"] == 0 and d["disputes"] == {"open": [], "resolved": []}
    kinds = [b["kind"] for b in d["briefing"]]
    assert kinds[0] == "call" and "review" in kinds and "contact" in kinds


def test_weighted_average_age_is_exact(server):
    asof = date(2026, 10, 15)
    owed = [(100000, date(2026, 7, 1)), (40000, date(2026, 9, 1)), (15000, date(2026, 9, 1)), (7000, date(2026, 10, 1))]   # (outstanding, due date)
    num = sum(a * (asof - due).days for a, due in owed)
    den = sum(a for a, _ in owed)
    assert overview(server)["kpi"]["avg_days_overdue"] == f"{num / den:.1f}"


def test_near_duplicate_customer_names_are_reported(server, tmp_path):
    with Store(server.db_path, now=lambda: NOW) as s:
        s.apply_import(import_rows([(1, HEADER), (2, ["INV-9", "ACME TRADERS PVT LTD", "01-09-2026", "01-10-2026", "100", "0"])]), actor="t")
    assert overview(server)["health"]["near_duplicate_names"] == [["ACME TRADERS PVT LTD", "Acme Traders"]]


def test_missing_due_dates_show_up_in_health_and_briefing(server):
    with Store(server.db_path, now=lambda: NOW) as s:
        s.apply_import(import_rows([(1, HEADER), (2, ["INV-7", "Zed Co", "01-09-2026", "", "900", "0"])]), actor="t")
    d = overview(server)
    assert d["health"]["no_due_date"] == {"count": 1, "amount": "900.00"}
    assert any(b["kind"] == "health" for b in d["briefing"])


# ---- customer detail, activity, backup -----------------------------------------------------------------


def test_customer_detail_shape(server):
    status, _, d = call(server, f"/api/customer?name=acme%20traders&{ASOF}")
    assert status == 200 and d["customer"] == "Acme Traders"
    assert d["summary"]["outstanding"] == "140000.00" and d["summary"]["overdue"] == "140000.00" and d["summary"]["oldest_days"] == 106
    assert [i["invoice_no"] for i in d["invoices"]] == ["INV-1001", "INV-1002"] and d["invoices"][0]["status"] == "overdue"
    assert d["stage"] == 4 and d["queue"]["customer"] == "Acme Traders" and d["reliability"] == {"kept": 0, "broken": 0}


def test_customer_detail_errors(server):
    assert call(server, "/api/customer?name=Nobody")[0] == 404
    assert call(server, "/api/customer")[0] == 400


def test_activity_feed(server):
    post(server, "/api/followup", {"customer": "Acme Traders", "channel": "call", "outcome": "x", "actor": "priya"})
    status, _, d = call(server, "/api/activity?limit=2")
    assert status == 200 and len(d["entries"]) == 2 and d["entries"][0]["action"] == "followup.logged" and d["entries"][0]["actor"] == "priya"
    older = call(server, f"/api/activity?limit=2&before={d['entries'][-1]['id']}")[2]["entries"]
    assert older and older[0]["id"] < d["entries"][-1]["id"]
    assert call(server, "/api/activity?limit=abc")[0] == 400


def test_backup_is_a_valid_database_with_the_data(server, tmp_path):
    status, headers, data = call(server, "/backup.db")
    assert status == 200 and "attachment" in headers["content-disposition"] and ".db" in headers["content-disposition"]
    out = tmp_path / "b.db"
    out.write_bytes(data)
    with sqlite3.connect(out) as db:
        assert db.execute("SELECT COUNT(*) FROM invoices").fetchone()[0] == 4


# ---- drafts with links -----------------------------------------------------------------------------------


def test_drafts_in_both_languages_with_links_when_contact_exists(server):
    post(server, "/api/customer/update", {"customer": "Acme Traders", "phone": "98765 43210", "email": "acct@acme.in"})
    _, _, en = call(server, f"/api/draft?customer=Acme%20Traders&{ASOF}")
    _, _, hi = call(server, f"/api/draft?customer=Acme%20Traders&language=hinglish&{ASOF}")
    assert en["language"] == "en" and hi["language"] == "hinglish" and "Namaste" in hi["email"] and "Namaste" not in en["email"]
    assert en["links"]["whatsapp"].startswith("https://wa.me/919876543210?text=") and en["links"]["tel"] == "tel:+919876543210"
    assert en["links"]["email"].startswith("mailto:acct@acme.in?subject=") and en["contact"]["phone"] == "+919876543210"
    assert call(server, f"/api/draft?customer=Acme&language=klingon&{ASOF}")[0] == 400


def test_no_links_without_contact_details_and_none_when_do_not_contact(server):
    _, _, d = call(server, f"/api/draft?customer=Acme%20Traders&{ASOF}")
    assert d["links"] == {"whatsapp": None, "email": None, "tel": None}
    post(server, "/api/customer/update", {"customer": "Acme Traders", "phone": "9876543210", "do_not_contact": True})
    _, _, d = call(server, f"/api/draft?customer=Acme%20Traders&{ASOF}")
    assert d["links"] is None and d["contact"]["do_not_contact"] is True


# ---- customer master, disputes, settings, settle -----------------------------------------------------------


def test_customer_update_effects_and_validation(server):
    s, _, out = post(server, "/api/customer/update", {"customer": "beta co", "phone": "9876543210", "credit_limit": "10,000", "notes": "pays late"})
    assert s == 200 and out["phone"] == "+919876543210" and out["credit_limit"] == "10000.00"
    beta = next(q for q in overview(server)["queue"] if q["customer"] == "Beta Co")
    assert beta["over_limit"] is True and beta["contact"]["phone"] == "+919876543210"
    assert post(server, "/api/customer/update", {"customer": "Beta Co", "phone": "abc"})[0] == 400
    assert post(server, "/api/customer/update", {"customer": "Nobody", "phone": "9876543210"})[0] == 400
    assert post(server, "/api/customer/update", {"customer": "Beta Co", "do_not_contact": "yes"})[0] == 400
    s, _, out = post(server, "/api/customer/update", {"customer": "Beta Co", "credit_limit": ""})
    assert s == 200 and out["credit_limit"] is None


def test_do_not_contact_moves_the_customer_to_waiting(server):
    post(server, "/api/customer/update", {"customer": "Gamma Ltd", "do_not_contact": True})
    d = overview(server)
    assert "Gamma Ltd" not in [q["customer"] for q in d["queue"]]
    assert any(w["customer"] == "Gamma Ltd" and "do-not-contact" in w["reason"] for w in d["waiting"])


def test_dispute_lifecycle_through_the_api(server):
    s, _, out = post(server, "/api/dispute/open", {"invoice_no": "INV-1001", "customer": "Acme Traders", "category": "price", "note": "rate", "owner": "Rohit"})
    assert s == 200
    d = overview(server)
    acme = next(q for q in d["queue"] if q["customer"] == "Acme Traders")
    assert acme["overdue"] == "40000.00" and acme["disputed_overdue"] == "100000.00"       # the disputed invoice is out of the chase
    assert d["disputes"]["open"][0]["category"] == "price" and d["kpi"]["disputed_count"] == 1
    assert post(server, "/api/dispute/open", {"invoice_no": "INV-1001", "customer": "Acme Traders", "category": "price"})[0] == 400
    assert post(server, "/api/dispute/open", {"invoice_no": "INV-1001", "customer": "Acme Traders", "category": "rudeness"})[0] == 400
    assert post(server, "/api/dispute/resolve", {"id": out["id"], "note": "credit note"})[0] == 200
    d = overview(server)
    assert next(q for q in d["queue"] if q["customer"] == "Acme Traders")["overdue"] == "140000.00" and d["disputes"]["resolved"][0]["resolution_note"] == "credit note"
    assert post(server, "/api/dispute/resolve", {"id": out["id"]})[0] == 400
    assert post(server, "/api/dispute/resolve", {"id": "1"})[0] == 400


def test_msme_setting_switches_the_informational_flag_on(server):
    assert overview(server)["kpi"]["msme_past_window"] == 0
    # three invoices are more than 45 days old; Gamma's (1 Sep, 44 days) is not
    assert post(server, "/api/setting", {"key": "msme_supplier", "value": "1"})[0] == 200
    d = overview(server)
    assert d["settings"]["msme_supplier"] is True and d["kpi"]["msme_past_window"] == 3 and "CA" in d["settings"]["msme_caution"]
    assert any(i["past_msme_window"] for i in call(server, f"/api/customer?name=Acme%20Traders&{ASOF}")[2]["invoices"])
    assert post(server, "/api/setting", {"key": "msme_supplier", "value": "maybe"})[0] == 400
    assert post(server, "/api/setting", {"key": "bogus", "value": "1"})[0] == 400


def test_settle_confirmed_invoices(server):
    s, _, out = post(server, "/api/invoice/settle", {"items": [{"invoice_no": "INV-2001", "customer": "Beta Co"}], "note": "NEFT seen"})
    assert s == 200 and out == {"settled": 1, "total": "15000.00"}
    assert "Beta Co" not in [q["customer"] for q in overview(server)["queue"]]
    assert post(server, "/api/invoice/settle", {"items": [{"invoice_no": "INV-2001", "customer": "Beta Co"}]})[0] == 400
    assert post(server, "/api/invoice/settle", {"items": "nope"})[0] == 400
    assert post(server, "/api/invoice/settle", {"items": []})[0] == 400


# ---- import preview / apply -------------------------------------------------------------------------------


def test_import_preview_describes_changes_without_saving(server):
    changed = ["INV-1001", "Acme Traders", "01-06-2026", "01-07-2026", "100000", "30000"]       # +30,000 paid
    new = ["INV-4001", "Delta", "01-10-2026", "01-11-2026", "5000", "0"]
    s, _, p = upload(server, "/api/import/preview", csv_bytes(changed, ROW2, ROWB, new, ["", "Grand Total", "", "", "", ""]))
    assert s == 200 and p["valid"] == 4 and p["error_count"] == 0 and p["skipped_rows"] == 1
    assert p["diff"]["new_count"] == 1 and p["diff"]["changed_count"] == 1 and p["diff"]["unchanged"] == 2
    assert p["diff"]["payments_seen"] == "30000.00" and p["diff"]["changed"][0]["delta"] == "30000.00"
    assert [m["invoice_no"] for m in p["diff"]["missing"]] == ["INV-3001"] and p["diff"]["missing_total"] == "7000.00"
    with Store(server.db_path) as st:                                        # nothing was written
        assert len(st.invoices()) == 4 and next(i for i in st.invoices() if i.invoice_no == "INV-1001").paid == 0
    assert any(server.uploads.directory.iterdir())                           # the file waits for confirmation


def test_import_preview_reports_row_errors_and_names_the_sheet(server):
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Cover"); ws.append(["notes"])
    ws = wb.create_sheet("Bills")
    ws.append(HEADER); ws.append(["INV-5", "Epsilon", "01-09-2026", "01-10-2026", "abc", "0"]); ws.append(["INV-6", "Epsilon", "01-09-2026", "01-10-2026", "500", "0"])
    buf = io.BytesIO(); wb.save(buf)
    s, _, p = upload(server, "/api/import/preview", buf.getvalue(), filename="weekly.xlsx")
    assert s == 200 and p["sheet"] == "Bills" and p["filename"] == "weekly.xlsx" and p["valid"] == 1
    assert p["error_count"] == 1 and p["errors"][0]["row"] == 2 and "not a valid amount" in p["errors"][0]["message"]


def test_import_apply_saves_snapshots_and_reports_what_to_confirm(server):
    changed = ["INV-1001", "Acme Traders", "01-06-2026", "01-07-2026", "100000", "30000"]
    _, _, p = upload(server, "/api/import/preview", csv_bytes(changed, ROW2, ROWB))
    s, _, out = post(server, "/api/import/apply", {"upload_id": p["upload_id"], "take_snapshot": True, "actor": "priya"})
    assert s == 200 and (out["inserted"], out["updated"], out["unchanged"]) == (0, 1, 2) and out["payments_seen"] == "30000.00"
    assert out["snapshot"]["as_of"] == date.today().isoformat() and [m["invoice_no"] for m in out["missing"]] == ["INV-3001"]
    with Store(server.db_path) as st:
        assert next(i for i in st.invoices() if i.invoice_no == "INV-1001").paid == 30000
        assert [e for e in st.audit_entries() if e.action == "invoice.updated"][0].actor == "priya"
    assert not any(server.uploads.directory.iterdir())                       # the temporary file is gone
    assert post(server, "/api/import/apply", {"upload_id": p["upload_id"]})[0] == 400      # cannot be applied twice


def test_apply_refuses_files_with_errors_unless_allowed(server):
    bad = ["INV-9", "Zed", "01-09-2026", "01-10-2026", "abc", "0"]
    _, _, p = upload(server, "/api/import/preview", csv_bytes(ROW1, bad))
    s, _, out = post(server, "/api/import/apply", {"upload_id": p["upload_id"]})
    assert s == 400 and "error row" in out["error"]
    assert post(server, "/api/import/apply", {"upload_id": p["upload_id"], "allow_errors": True})[0] == 200


def test_credit_days_fill_missing_due_dates_only_when_asked(server):
    row = ["INV-8", "Eta Co", "01-09-2026", "", "900", "0"]
    _, _, plain = upload(server, "/api/import/preview", csv_bytes(row))
    assert plain["filled_due_dates"] == 0 and any("no due date" in w for w in plain["warnings"])
    s, _, filled = call(server, "/api/import/preview?credit_days=30", "POST",
                        headers={"Content-Type": "application/octet-stream", web.TOKEN_HEADER: token(server), "X-Filename": "a.csv"}, raw=csv_bytes(row))
    assert s == 200 and filled["filled_due_dates"] == 1
    assert upload(server, "/api/import/preview?credit_days=abc", csv_bytes(row))[0] in (200, 400)


@pytest.mark.parametrize("data, name, fragment", [
    (b"PK\x03\x04junk", "x.xlsx", "damaged"),
    (b"data", "x.exe", ".csv or .xlsx"),
    (b"", "x.csv", "empty"),
])
def test_bad_uploads_get_clear_400s_and_leave_no_files(server, data, name, fragment):
    s, _, out = upload(server, "/api/import/preview", data, filename=name)
    assert s == 400 and fragment in out["error"].lower()
    assert not any(server.uploads.directory.iterdir())


def test_uploaded_filenames_cannot_escape_or_be_used_on_disk(server):
    s, _, p = upload(server, "/api/import/preview", csv_bytes(ROW1), filename="../../../../tmp/evil.csv")
    assert s == 200 and p["filename"] == "evil.csv"
    assert all(f.parent == server.uploads.directory for f in server.uploads.directory.iterdir())


def test_uploads_are_deleted_when_the_server_stops(tmp_path):
    srv = web.make_server(make_db(tmp_path))
    threading.Thread(target=lambda: srv.serve_forever(poll_interval=0.02), daemon=True).start()
    upload(srv, "/api/import/preview", csv_bytes(ROW1))
    directory = srv.uploads.directory
    assert any(directory.iterdir())
    srv.shutdown(); srv.server_close()
    assert not directory.exists()


# ---- bank statement matching ------------------------------------------------------------------------------


STATEMENT = (b"Date,Narration,Chq./Ref.No.,Withdrawal Amt.,Deposit Amt.\n"
             b"01-10-2026,NEFT CR-ACME TRADERS-INV-1001,N1,,100000.00\n"
             b"02-10-2026,UPI/BETA CO,U2,,14700.00\n"
             b"03-10-2026,ATM,A3,500.00,\n"
             b"04-10-2026,UPI RANDOM 77,U4,,123.45\n")


def test_statement_matching_returns_explained_suggestions_and_changes_nothing(server):
    s, _, out = upload(server, "/api/statement/match", STATEMENT, filename="bank.csv")
    assert s == 200 and out["credits"] == 3 and out["skipped"] == 1
    first, second, third = out["suggestions"]
    assert first["confidence"] == "high" and first["customer"] == "Acme Traders" and first["invoice_nos"] == ["INV-1001"]
    assert second["customer"] == "Beta Co" and any("TDS" in r for r in second["reasons"])        # 14,700 is 2% below 15,000
    assert third["confidence"] == "none"
    with Store(server.db_path) as st:
        assert st.invoices()[0].paid == 0 and len(st.audit_entries()) == 4                         # unchanged
    assert not any(server.uploads.directory.iterdir())


def test_statement_with_the_wrong_columns_is_a_clean_400(server):
    s, _, out = upload(server, "/api/statement/match", csv_bytes(ROW1), filename="inv.csv")
    assert s == 400 and "statement columns" in out["error"]


# ---- every new write route is guarded -----------------------------------------------------------------------

NEW_ROUTES = [r for r in web.WRITE_ROUTES]


@pytest.mark.parametrize("path", NEW_ROUTES)
def test_every_write_route_demands_the_token_origin_and_content_type(server, path):
    kind = web.WRITE_ROUTES[path][0]
    good_type = web.CONTENT_TYPES[kind]
    body = b"{}" if kind == "json" else b"x"
    base = {"Content-Type": good_type}
    assert call(server, path, "POST", headers=base, raw=body)[0] == 403                                           # no token
    assert call(server, path, "POST", headers={**base, web.TOKEN_HEADER: "wrong"}, raw=body)[0] == 403
    tok = {web.TOKEN_HEADER: token(server)}
    assert call(server, path, "POST", headers={**base, **tok, "Origin": "http://evil.example"}, raw=body)[0] == 403
    assert call(server, path, "POST", headers={"Content-Type": "text/plain", **tok}, raw=body)[0] == 415
    assert call(server, path, "POST", headers={**base, **tok}, raw=body, host="evil.example")[0] == 403
    other = "application/octet-stream" if kind == "json" else "application/json"
    assert call(server, path, "POST", headers={"Content-Type": other, **tok}, raw=body)[0] == 415                 # the other kind's type is refused


@pytest.mark.parametrize("path", NEW_ROUTES)
def test_read_only_mode_blocks_every_write_route(tmp_path, path):
    srv = web.make_server(make_db(tmp_path), read_only=True)
    threading.Thread(target=lambda: srv.serve_forever(poll_interval=0.02), daemon=True).start()
    try:
        kind = web.WRITE_ROUTES[path][0]
        assert call(srv, path, "POST", headers={"Content-Type": web.CONTENT_TYPES[kind], web.TOKEN_HEADER: srv.token}, raw=b"{}" if kind == "json" else b"x")[0] == 405
    finally:
        srv.shutdown(); srv.server_close()


def test_upload_size_limit(server, monkeypatch):
    big = b"a" * (web.MAX_UPLOAD + 1)
    try:
        s, _, _ = upload(server, "/api/import/preview", big)
        assert s == 413
    except (ConnectionResetError, BrokenPipeError):
        pass


def test_a_file_with_unrecognised_columns_is_kept_for_mapping_not_rejected(server):
    s, _, p = upload(server, "/api/import/preview", b"not a spreadsheet at all\nsecond line\n", filename="odd.csv")
    assert s == 200 and p["needs_mapping"] is True and p["filename"] == "odd.csv" and "could not find an invoice table" in p["message"]
    assert p["ai_available"] is False and len(list(server.uploads.directory.iterdir())) == 1       # waiting for the person
