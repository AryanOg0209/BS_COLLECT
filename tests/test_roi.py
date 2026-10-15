from decimal import Decimal

import pytest
from openpyxl import load_workbook

from collectos.cli import main
from collectos.roi import (
    VERDICT_NEGATIVE,
    VERDICT_POSITIVE,
    RoiInputs,
    calculate,
    financing_saving_at,
    labour_saving_at,
    write_calculator,
)

D = Decimal


def test_default_numbers_worked_by_hand():
    # 10 h/week * 52/12 = 43.3333 h + 6 h reporting = 49.3333 h/month
    r = calculate(RoiInputs())
    assert r.hours_per_month == D("49.33")
    assert r.hours_saved_per_month == D("9.87")             # 20% of 49.3333
    assert r.labour_saving_per_month == D("1973.33")        # 9.8667 h * Rs 200
    assert r.net_labour_after_fee == D("-526.67")           # 1973.33 - 2500
    assert r.pilot_payback_months == D("2.03")              # 4000 / 1973.33
    assert r.financing_saving_per_month == D("0.00")        # no days-faster assumed
    assert r.verdict == VERDICT_NEGATIVE


def test_positive_case_and_verdict():
    r = calculate(RoiInputs(collection_hours_per_week=D("25"), monthly_fee=D("2000")))
    assert r.net_labour_after_fee > 0 and r.verdict == VERDICT_POSITIVE


def test_zero_reduction_gives_no_payback_not_a_crash():
    r = calculate(RoiInputs(reduction_pct=D("0")))
    assert r.labour_saving_per_month == D("0.00") and r.pilot_payback_months is None


def test_financing_is_separate_and_exact():
    # Rs 25,00,000 a month, 5 days faster, 12% a year: 2,500,000 * 5/365 * 0.12
    i = RoiInputs(days_faster=D("5"))
    assert calculate(i).financing_saving_per_month == D("4109.59")
    assert calculate(i).labour_saving_per_month == calculate(RoiInputs()).labour_saving_per_month  # unaffected
    assert financing_saving_at(RoiInputs(), D("10")) == D("8219.18")


def test_sensitivity_helpers():
    i = RoiInputs()
    assert labour_saving_at(i, D("10")) == D("986.67") and labour_saving_at(i, D("40")) == D("3946.67")


def test_workbook_has_live_formulas_and_editable_inputs(tmp_path):
    ws = load_workbook(write_calculator(tmp_path / "roi.xlsx"))["Calculator"]
    formulas = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) == 19  # 5 outputs, reading, financing, 8 sensitivity, 4 scenarios
    assert ws["B7"].value == 10 and ws["B7"].fill.fgColor.rgb.endswith("FFF2CC")  # hours/week input
    assert ws.data_validations.dataValidation  # numeric guard on inputs
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "not predictions or guarantees" in text and "not recoverable money" in text


def test_workbook_reflects_custom_defaults(tmp_path):
    ws = load_workbook(write_calculator(tmp_path / "roi.xlsx", RoiInputs(hourly_cost=D("350"))))["Calculator"]
    assert ws["B9"].value == 350


def test_cli_roi_sheet(capsys, tmp_path):
    out = tmp_path / "r.xlsx"
    assert main(["roi-sheet", str(out)]) == 0 and out.exists()
    assert "illustrative placeholders" in capsys.readouterr().out
