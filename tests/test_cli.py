from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from collectos import synthetic
from collectos.cli import main
from collectos.importer import import_invoices, import_rows
from collectos.store import Store

HEADER = "Invoice No,Customer,Invoice Date,Due Date,Amount,Paid\n"


def d(offset_days):
    return (date.today() + timedelta(days=offset_days)).strftime("%d-%m-%Y")


@pytest.fixture
def db(tmp_path):
    return str(tmp_path / "t.db")


@pytest.fixture
def csv_file(tmp_path):
    path = tmp_path / "inv.csv"
    path.write_text(
        HEADER
        + f"A1,Acme Traders,{d(-120)},{d(-90)},100000,0\n"
        + f"A2,Acme Traders,{d(-40)},{d(-10)},5000,0\n"
        + f"B1,Beta Co,{d(-20)},{d(10)},2000,0\n",
        encoding="utf-8",
    )
    return str(path)


def run(capsys, *argv):
    code = main(list(argv))
    out = capsys.readouterr()
    return code, out.out, out.err


def test_import_then_reimport_is_idempotent(capsys, db, csv_file):
    code, out, _ = run(capsys, "import", csv_file, "--db", db, "--actor", "t")
    assert code == 0 and "3 new, 0 changed, 0 unchanged" in out
    code, out, _ = run(capsys, "import", csv_file, "--db", db, "--actor", "t")
    assert code == 0 and "0 new, 0 changed, 3 unchanged" in out


def test_import_with_errors_saves_nothing_unless_allowed(capsys, db, tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text(HEADER + f"A1,Acme,{d(-50)},{d(-20)},100,0\nX1,Bad,{d(-50)},{d(-20)},abc,0\n", encoding="utf-8")
    code, out, _ = run(capsys, "import", str(bad), "--db", db)
    assert code == 2 and "row 3: not a valid amount" in out and "Nothing saved" in out
    with Store(db) as s:
        assert s.invoices() == []
    code, out, _ = run(capsys, "import", str(bad), "--db", db, "--allow-errors")
    assert code == 0 and "1 new" in out


def test_queue_log_and_resolve_flow(capsys, db, csv_file):
    run(capsys, "import", csv_file, "--db", db)
    code, out, _ = run(capsys, "queue", "--db", db)
    assert code == 0 and "Acme Traders" in out and "Beta Co" not in out  # Beta is not yet due
    assert "never contacted" in out

    future = (date.today() + timedelta(days=5)).strftime("%d-%m-%Y")
    code, out, _ = run(capsys, "log", "acme traders", "--db", db, "--channel", "call", "--outcome", "Will pay",
                       "--promise-date", future, "--promise-amount", "50,000", "--actor", "priya")
    assert code == 0 and "promise #1" in out

    code, out, _ = run(capsys, "queue", "--db", db)
    assert "Waiting" in out and "promised to pay by" in out

    code, out, _ = run(capsys, "resolve", "1", "kept", "--db", db, "--note", "NEFT", "--actor", "priya")
    assert code == 0 and "marked kept" in out
    code, out, err = run(capsys, "resolve", "1", "broken", "--db", db)
    assert code == 1 and "already kept" in err


def test_user_errors_exit_1_with_a_message_not_a_traceback(capsys, db, csv_file):
    run(capsys, "import", csv_file, "--db", db)
    code, _, err = run(capsys, "log", "Nobody Ltd", "--db", db, "--channel", "call", "--outcome", "x")
    assert code == 1 and err.startswith("error: unknown customer")
    code, _, err = run(capsys, "import", "/no/such/file.csv", "--db", db)
    assert code == 1 and err.startswith("error:")
    code, _, err = run(capsys, "log", "Acme Traders", "--db", db, "--channel", "call", "--outcome", "x", "--promise-date", "not-a-date")
    assert code == 1 and "unrecognised date" in err


def test_draft_prints_a_reminder_and_never_sends(capsys, db, csv_file):
    run(capsys, "import", csv_file, "--db", db)
    code, out, _ = run(capsys, "draft", "Acme Traders", "--db", db, "--sender", "Rohit", "--company", "Sharma Steel")
    assert code == 0
    assert "[template: firm]" in out and "Subject: Overdue payment: 2 invoices, ₹1,05,000.00" in out and "A1" in out and "WhatsApp version" in out
    code, _, err = run(capsys, "draft", "Beta Co", "--db", db, "--sender", "R", "--company", "C")
    assert code == 1 and "no overdue invoices" in err  # Beta is not yet due


def test_report_command_writes_a_workbook(capsys, db, csv_file, tmp_path):
    run(capsys, "import", csv_file, "--db", db)
    out_file = tmp_path / "r.xlsx"
    code, out, _ = run(capsys, "report", str(out_file), "--db", db, "--company", "Test Co")
    assert code == 0 and out_file.exists()
    assert "Summary" in load_workbook(out_file).sheetnames


# ---- synthetic data & demo --------------------------------------------------------


def test_synthetic_data_is_deterministic_and_has_the_planted_problems(tmp_path):
    as_of = date(2026, 10, 15)
    assert synthetic.generate_rows(as_of) == synthetic.generate_rows(as_of)
    assert synthetic.generate_rows(as_of, seed=1) != synthetic.generate_rows(as_of, seed=2)

    report = import_rows([(i, r) for i, r in enumerate(synthetic.generate_rows(as_of), start=1)], source="synthetic")
    assert len(report.errors) == 2 and len(report.duplicates) == 2
    assert {e.message.split(":")[0] for e in report.errors} == {"not a valid amount", "unrecognised date format"}
    assert sum(1 for i in report.invoices if i.invoice.due_date is None) == 3
    assert len(report.skipped_rows) == 1  # Grand Total
    assert all(i.invoice.customer.endswith("(Demo)") for i in report.invoices)
    assert all(isinstance(i.invoice.amount, Decimal) for i in report.invoices)


def test_demo_builds_the_full_kit(capsys, tmp_path):
    out_dir = tmp_path / "kit"
    code, out, _ = run(capsys, "demo", str(out_dir), "--as-of", "15-10-2026")
    assert code == 0
    for name in ("synthetic_invoices.csv", "owner_report.xlsx", "reminder_drafts.txt", "demo.db"):
        assert (out_dir / name).stat().st_size > 0
    assert "broken promise (was due 05-Oct-2026)" in out          # scripted behaviours all visible
    assert "promised to pay by 19-Oct-2026" in out
    assert "contacted 1 day(s) ago" in out
    drafts = (out_dir / "reminder_drafts.txt").read_text(encoding="utf-8")
    assert "[promise_followup]" in drafts and "Your Name" in drafts
    wb = load_workbook(out_dir / "owner_report.xlsx")
    assert wb["Summary"]["A1"].value.startswith("Receivables summary: Demo Distributors (synthetic data)")
    with Store(out_dir / "demo.db") as s:
        assert len(s.open_promises()) == 2 and len(s.followups()) == 5
    # Re-running into a non-empty folder is refused, never overwritten.
    code, _, err = run(capsys, "demo", str(out_dir))
    assert code == 1 and "not empty" in err


def test_demo_is_reproducible(capsys, tmp_path):
    for name in ("a", "b"):
        run(capsys, "demo", str(tmp_path / name), "--as-of", "15-10-2026")
    assert (tmp_path / "a" / "synthetic_invoices.csv").read_text() == (tmp_path / "b" / "synthetic_invoices.csv").read_text()
    assert (tmp_path / "a" / "reminder_drafts.txt").read_text() == (tmp_path / "b" / "reminder_drafts.txt").read_text()


# ---- snapshots, trend, history ------------------------------------------------------


def test_snapshot_and_trend_commands(capsys, db, csv_file):
    run(capsys, "import", csv_file, "--db", db)
    code, out, _ = run(capsys, "trend", "--db", db)
    assert code == 0 and "No snapshots yet" in out
    code, out, _ = run(capsys, "snapshot", "--db", db)
    assert code == 0 and out.startswith("Baseline snapshot") and "₹1,05,000.00" in out
    code, out, _ = run(capsys, "snapshot", "--db", db)
    assert out.startswith("Snapshot for")
    code, out, _ = run(capsys, "trend", "--db", db)
    assert "Paid seen" in out and out.count("\n") >= 2


def test_report_command_includes_trend_when_snapshots_exist(capsys, db, csv_file, tmp_path):
    run(capsys, "import", csv_file, "--db", db)
    run(capsys, "snapshot", "--db", db)
    out_file = tmp_path / "r.xlsx"
    run(capsys, "report", str(out_file), "--db", db, "--company", "X")
    assert "Trend" in load_workbook(out_file).sheetnames


def test_synthetic_history_is_consistent_with_the_final_export():
    as_of = date(2026, 10, 15)
    final = {i.invoice_no: i for i in synthetic.generate_invoices(as_of)}
    history = synthetic.generate_history(as_of)
    assert sorted(history) == [7, 14, 21, 28]
    previous_paid = {}
    for ago in (28, 21, 14, 7):
        day = as_of - timedelta(days=ago)
        report = import_rows([(i, r) for i, r in enumerate(history[ago], start=1)])
        assert report.ok
        for item in report.invoices:
            inv, truth = item.invoice, final[item.invoice.invoice_no]
            assert inv.invoice_date <= day                       # no invoice exists before it is issued
            assert inv.paid in (Decimal("0.00"), truth.final_paid)
            assert inv.paid >= previous_paid.get(inv.invoice_no, Decimal("0.00"))  # payments only accumulate
            previous_paid[inv.invoice_no] = inv.paid
    final_paid_seen = {n: p for n, p in previous_paid.items()}
    assert all(p <= final[n].final_paid for n, p in final_paid_seen.items())


def test_demo_records_five_weekly_snapshots_with_payments_seen(capsys, tmp_path):
    code, out, _ = run(capsys, "demo", str(tmp_path / "kit"), "--as-of", "15-10-2026")
    assert code == 0 and "Synthetic five-week history" in out
    with Store(tmp_path / "kit" / "demo.db") as s:
        snaps = s.snapshots()
    assert [x.as_of for x in snaps] == [date(2026, 9, 17) + timedelta(days=7 * k) for k in range(5)]
    assert snaps[0].payments_observed == Decimal("0.00")             # baseline
    assert all(x.payments_observed > 0 for x in snaps[1:])
    wb = load_workbook(tmp_path / "kit" / "owner_report.xlsx")
    assert "Trend" in wb.sheetnames and wb["Trend"].max_row >= 10
