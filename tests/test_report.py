from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from collectos.models import Invoice, Promise, customer_key
from collectos.report import write_owner_report
from collectos.store import FollowUp
from collectos.workqueue import build_queue

TODAY = date(2026, 10, 15)


def days_ago(n):
    return date.fromordinal(TODAY.toordinal() - n)


def inv(no, customer, due_in_days_ago, amount="10000", paid="0"):
    return Invoice(no, customer, date(2025, 1, 1), days_ago(due_in_days_ago), Decimal(amount), Decimal(paid))


def make(tmp_path, invoices, promises=(), followups=(), last=None, warnings=()):
    queue = build_queue(invoices, TODAY, last or {}, promises)
    path = write_owner_report(
        tmp_path / "report.xlsx", company="Demo Co", as_of=TODAY, invoices=invoices, queue=queue,
        open_promises=promises, followups=followups, warnings=warnings,
    )
    return load_workbook(path)


def column(ws, col):
    return [ws.cell(row=r, column=col).value for r in range(1, ws.max_row + 1)]


def dec(cell_value):
    return Decimal(str(cell_value))


SCENARIO = [
    inv("A1", "Acme", 100, "200000"),               # 91+ , big
    inv("A2", "Acme", -10, "5000"),                 # not due
    inv("B1", "Beta Traders", 40, "30000", "10000"),  # part-paid, 31-60
    inv("C1", "Cee", 5, "999.99"),                  # 1-30
    inv("D1", "Dee", 50, "100", "100"),             # settled, excluded
]


def test_sheets_present(tmp_path):
    wb = make(tmp_path, SCENARIO)
    assert wb.sheetnames == ["Summary", "Today's Queue", "Waiting", "Ageing by Customer", "Open Invoices", "Promises", "Notes"]


def test_summary_numbers_are_exact(tmp_path):
    ws = make(tmp_path, SCENARIO)["Summary"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 20)}
    assert dec(values["Total outstanding"]) == Decimal("225999.99")  # 200000 + 5000 + 20000 + 999.99
    assert dec(values["Overdue"]) == Decimal("220999.99")
    assert values["Overdue share of outstanding"] == "98%"
    assert values["Open invoices"] == 4
    assert values["Customers with overdue balance"] == 3


def test_ageing_table_on_summary(tmp_path):
    ws = make(tmp_path, SCENARIO)["Summary"]
    rows = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value) for r in range(1, 40)}
    assert rows["Current"][0] == 1 and dec(rows["Current"][1]) == Decimal("5000.00")
    assert rows["91+"][0] == 1 and dec(rows["91+"][1]) == Decimal("200000.00")
    assert rows["31-60"][0] == 1 and dec(rows["31-60"][1]) == Decimal("20000.00")
    assert rows["61-90"][0] == 0


def test_queue_sheet_lists_reasons_in_priority_order(tmp_path):
    ws = make(tmp_path, SCENARIO)["Today's Queue"]
    assert [ws.cell(row=r, column=2).value for r in range(2, 5)] == ["Acme", "Beta Traders", "Cee"]
    assert "overdue: +30" in ws.cell(row=2, column=9).value


def test_open_invoices_exclude_settled_and_include_partials(tmp_path):
    ws = make(tmp_path, SCENARIO)["Open Invoices"]
    numbers = column(ws, 1)[1:]
    assert "D1" not in numbers and "B1" in numbers
    b1 = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "B1")
    assert ws.cell(row=b1, column=11).value == "yes" and dec(ws.cell(row=b1, column=7).value) == Decimal("20000.00")


def test_indian_number_format_and_date_format_applied(tmp_path):
    ws = make(tmp_path, SCENARIO)["Open Invoices"]
    assert "##\\,##\\,##0.00" in ws["E2"].number_format
    assert ws["C2"].number_format == "DD-MMM-YYYY"


def test_formula_injection_is_neutralised(tmp_path):
    evil = [inv("X1", '=HYPERLINK("http://evil.example","click")', 10), inv("X2", "+cmd|calc", 10), inv("=1+1", "Normal", 10)]
    wb = make(tmp_path, evil)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                    assert cell.data_type == "s", f"{ws.title}!{cell.coordinate} would be parsed as a formula"
    found = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("HYPERLINK" in v for v in found)  # preserved as visible text, not dropped


def test_promises_and_followups_flow_into_the_report(tmp_path):
    promises = [
        Promise(1, "Acme", customer_key("Acme"), days_ago(2), Decimal("5000"), "open"),
        Promise(2, "Cee", customer_key("Cee"), date(2026, 10, 20), None, "open"),
    ]
    followups = [
        FollowUp(1, "Acme", "2026-10-14T10:00:00+00:00", "a", "call", "x", ""),
        FollowUp(2, "Acme", "2026-10-01T10:00:00+00:00", "a", "call", "old", ""),
    ]
    wb = make(tmp_path, SCENARIO, promises=promises, followups=followups)
    summary = {wb["Summary"].cell(row=r, column=1).value: wb["Summary"].cell(row=r, column=2).value for r in range(1, 20)}
    assert summary["Open promises past their date (broken)"] == 1
    assert summary["Follow-ups logged in the last 7 days"] == 1
    states = {wb["Promises"].cell(row=r, column=1).value: wb["Promises"].cell(row=r, column=4).value for r in (2, 3)}
    assert states == {"Acme": "BROKEN (date passed)", "Cee": "pending"}


def test_notes_carry_warnings_and_disclaimer(tmp_path):
    ws = make(tmp_path, SCENARIO, warnings=["2 invoice(s) have no due date and cannot be aged"])["Notes"]
    text = "\n".join(str(v) for v in column(ws, 1) if v)
    assert "no due date and cannot be aged" in text and "does not predict or guarantee" in text


def test_empty_input_still_produces_a_valid_report(tmp_path):
    wb = make(tmp_path, [])
    summary = {wb["Summary"].cell(row=r, column=1).value: wb["Summary"].cell(row=r, column=2).value for r in range(1, 20)}
    assert dec(summary["Total outstanding"]) == Decimal("0.00") and summary["Overdue share of outstanding"] is None
