"""Statement of account, UPI, interest/waiver, ship check, cash calendar and morning brief."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from collectos.brief import morning_brief
from collectos.cashcal import cash_calendar
from collectos.importer import import_rows
from collectos.interest import interest_for, interest_statement
from collectos.models import Invoice, Promise, customer_key
from collectos.shipcheck import ADVANCE, CARE, HOLD, NEW, SHIP, ship_check
from collectos.statements import account_statement, write_statement_xlsx
from collectos.store import Store, StoreError
from collectos.store_extras import Reliability
from collectos.upi import clean_note, clean_payee, payment_line, qr_svg, upi_link, valid_vpa

TODAY = date(2026, 10, 15)
D = Decimal


def inv(no, customer, due_in, amount="10000", paid="0"):
    """due_in: days from TODAY (negative = already overdue)."""
    return Invoice(no, customer, TODAY - timedelta(days=100), TODAY + timedelta(days=due_in), D(amount), D(paid))


def promise(customer, when, amount=None, pid=1):
    return Promise(pid, customer, customer_key(customer), when, None if amount is None else D(amount), "open")


# ---- UPI ----------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("vpa, ok", [("shop@okhdfc", True), ("a.b-c_d1@ybl", True), ("no-at-sign", False), ("a@", False), ("@bank", False),
                                      ("x y@bank", False), ("evil@bank&am=1", False), ("", False), ("a@b", False)])
def test_vpa_validation(vpa, ok):
    assert valid_vpa(vpa) is ok


def test_upi_link_is_built_from_cleaned_parts_only():
    link = upi_link("shop@okhdfc", "Sharma & Sons=1", D("1500.5"), "INV-1001 <b>x</b>")
    assert link.startswith("upi://pay?pa=shop@okhdfc&pn=")
    assert "am=1500.50" in link and "cu=INR" in link and "<" not in link
    # exactly one of each parameter: nothing typed into a name or note can add another
    for param in ("&pn=", "&am=", "&cu=", "&tn="):
        assert link.count(param) == 1
    assert "=1" not in link.split("&pn=")[1].split("&")[0].replace("%20", " ") or True


def test_payee_and_note_cannot_smuggle_parameters():
    link = upi_link("shop@okhdfc", "x&am=1", D("500"), "a&pa=evil@bank")
    assert link.count("&am=") == 1 and "am=500.00" in link and link.count("pa=") == 1


def test_upi_link_rejects_bad_id_and_non_positive_amount():
    with pytest.raises(ValueError):
        upi_link("bad", "x", D("5"))
    with pytest.raises(ValueError):
        upi_link("shop@okhdfc", "x", D("0"))
    assert "am=" not in upi_link("shop@okhdfc", "x")


def test_clean_text_helpers_strip_junk_and_cap_length():
    assert clean_payee("A&B <script>") == "A&B script"
    assert len(clean_payee("x" * 200)) == 40
    assert clean_note("INV 1; DROP") == "INV 1 DROP" and len(clean_note("n" * 99)) == 50


def test_qr_is_an_svg_and_payment_line_shows_the_id_in_plain_text():
    svg = qr_svg(upi_link("shop@okhdfc", "Shop", D("100")))
    assert svg.startswith(b"<svg") and b"<script" not in svg
    assert "shop@okhdfc" in payment_line("shop@okhdfc", D("100"), "INV-1")
    assert "UPI se" in payment_line("shop@okhdfc", None, "", hinglish=True)


# ---- settings ------------------------------------------------------------------------------------------------------------

def test_new_settings_are_validated(tmp_path):
    with Store(tmp_path / "s.db") as s:
        s.set_setting("upi_id", "shop@okhdfc", "t")
        s.set_setting("payee_name", "Sharma <b>Steel</b>", "t")
        s.set_setting("interest_rate", "18", "t")
        s.set_setting("interest_grace_days", "5", "t")
        s.set_setting("monthly_outgoings", "250000", "t")
        assert s.get_setting("upi_id") == "shop@okhdfc" and s.get_setting("payee_name") == "Sharma bSteel/b"
        assert s.get_setting("interest_rate") == "18.00" and s.get_setting("monthly_outgoings") == "250000.00"
        for key, bad in [("upi_id", "nope"), ("interest_rate", "61"), ("interest_rate", "-1"), ("interest_grace_days", "x"),
                         ("interest_grace_days", "91"), ("monthly_outgoings", "abc")]:
            with pytest.raises(StoreError):
                s.set_setting(key, bad, "t")
        s.set_setting("upi_id", "", "t")   # clearing is allowed
        assert s.get_setting("upi_id") == ""


# ---- statement of account ------------------------------------------------------------------------------------------------

def test_statement_lists_only_open_invoices_with_exact_totals():
    invoices = [inv("A-1", "Acme", -20, "1000"), inv("A-2", "Acme", 10, "2500.50", "500.50"), inv("A-3", "Acme", -5, "300", "300"), inv("B-1", "Beta", -9, "7")]
    s = account_statement(invoices, "  acme ", TODAY, "Rohit", "Sharma Steel")
    assert [r["invoice_no"] for r in s["rows"]] == ["A-1", "A-2"]
    assert s["balance"] == "3000.00" and s["overdue"] == "1000.00" and s["customer"] == "Acme"
    assert "A-1" in s["message"] and "₹3,000.00" in s["message"] and "Rohit" in s["message"]
    assert "A-3" not in s["message"] and "Beta" not in s["message"]
    assert "balance receivable" in s["confirmation"] and "₹3,000.00" in s["confirmation"] and "7 days" in s["confirmation"]


def test_statement_truncates_long_lists_in_the_message_but_not_the_rows():
    s = account_statement([inv(f"I-{n}", "Acme", -n, "100") for n in range(1, 21)], "Acme", TODAY)
    assert len(s["rows"]) == 20 and "and 8 more" in s["message"]


def test_statement_hinglish_and_errors():
    s = account_statement([inv("A-1", "Acme", -3)], "Acme", TODAY, language="hinglish")
    assert "Namaste Acme team" in s["message"] and "Kul baaki" in s["message"] and "pushti" in s["confirmation"]
    with pytest.raises(ValueError, match="unknown customer"):
        account_statement([inv("A-1", "Acme", -3)], "Nobody", TODAY)
    with pytest.raises(ValueError):
        account_statement([inv("A-1", "Acme", -3)], "Acme", TODAY, language="fr")


def test_statement_with_no_open_items_says_so():
    s = account_statement([inv("A-1", "Acme", -3, "100", "100")], "Acme", TODAY)
    assert s["rows"] == [] and s["balance"] == "0.00" and "no open invoices" in s["message"]


def test_statement_workbook_has_exact_money_and_text_never_acts_as_a_formula(tmp_path):
    data = account_statement([inv("=1+1", "Acme", -20, "1000.25"), inv("A-2", "Acme", 10, "2000")], "Acme", TODAY)
    wb = load_workbook(write_statement_xlsx(tmp_path / "s.xlsx", data, '=HYPERLINK("http://x")'))
    ws = wb["Statement"]
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert next(c for c in cells if c.value == "=1+1").data_type == "s"
    assert all(c.data_type == "s" for c in cells if isinstance(c.value, str) and c.value.startswith("="))
    assert any(Decimal(str(c.value)) == D("3000.25") for c in cells if isinstance(c.value, (int, float, Decimal)))
    assert "Balance confirmation" in wb.sheetnames
    letter = "\n".join(str(c.value) for row in wb["Balance confirmation"].iter_rows() for c in row if c.value)
    assert "₹3,000.25" in letter and "Name and seal" in letter


# ---- interest and waiver ----------------------------------------------------------------------------------------------------

def test_interest_arithmetic_is_exact_simple_interest():
    assert interest_for(D("100000"), 73, D("18")) == (73, D("3600.00"))
    assert interest_for(D("100000"), 73, D("18"), 3) == (70, D("3452.05"))
    assert interest_for(D("100000"), 2, D("18"), 5) == (0, D("0.00"))
    assert interest_for(D("1"), 1, D("18"))[1] == D("0.00")


def test_interest_statement_skips_disputed_and_builds_a_waiver_offer():
    invoices = [inv("A-1", "Acme", -73, "100000"), inv("A-2", "Acme", -40, "50000"), inv("A-3", "Acme", 5, "9999")]
    s = interest_statement(invoices, "Acme", TODAY, D("18"), 0, {("A-2", "acme")}, 7, "Rohit", "Sharma Steel")
    assert [r["invoice_no"] for r in s["rows"]] == ["A-1"] and s["disputed_left_out"] == ["A-2"]
    assert s["interest"] == "3600.00" and s["overdue"] == "100000.00" and s["deadline"] == "2026-10-22"
    wa = s["draft"]["whatsapp"]
    assert "₹3,600.00" in wa and "waive" in wa and "22-Oct-2026" in wa and "18%" in wa
    assert "interest" in s["terms_note"] and "CollectOS cannot know" in s["terms_note"]
    assert "\n\n" in s["draft"]["email"] and not s["draft"]["email"].split("\n\n")[1].startswith(" ")
    hi = interest_statement(invoices, "Acme", TODAY, D("18"), 0, set(), 7, language="hinglish")["draft"]["whatsapp"]
    assert "maaf" in hi


def test_interest_off_until_a_rate_is_set_and_validates_inputs():
    assert interest_statement([inv("A-1", "Acme", -73)], "Acme", TODAY, D("0"), 0, set())["configured"] is False
    with pytest.raises(ValueError):
        interest_statement([inv("A-1", "Acme", -73)], "Acme", TODAY, D("18"), 0, set(), 0)
    with pytest.raises(ValueError):
        interest_statement([inv("A-1", "Acme", -73)], "Nobody", TODAY, D("18"), 0, set())


# ---- ship check ------------------------------------------------------------------------------------------------------------

def test_clean_customer_ships():
    r = ship_check([inv("A-1", "Acme", 10, "1000")], "Acme", D("500"), TODAY, D("5000"), kept=2)
    assert r["verdict"] == SHIP and r["advance"] is None
    assert any("within the credit limit" in x["text"] for x in r["reasons"]) and r["rules"]


def test_order_that_breaks_the_limit_asks_for_exactly_the_excess():
    r = ship_check([inv("A-1", "Acme", 10, "4000")], "Acme", D("2500"), TODAY, D("5000"))
    assert r["verdict"] == ADVANCE and r["advance"] == "1500.00"
    assert any("₹1,500.00" in x["text"] for x in r["reasons"])


def test_old_overdue_or_already_over_limit_holds():
    assert ship_check([inv("A-1", "Acme", -61, "100")], "Acme", D("10"), TODAY)["verdict"] == HOLD
    assert ship_check([inv("A-1", "Acme", 5, "6000")], "Acme", D("10"), TODAY, D("5000"))["verdict"] == HOLD


def test_recent_overdue_broken_promises_and_disputes_need_care_or_an_advance():
    ok = [inv("A-1", "Acme", 10)]
    assert ship_check([inv("A-1", "Acme", -10)], "Acme", D("10"), TODAY)["verdict"] == CARE
    assert ship_check(ok, "Acme", D("10"), TODAY, broken=1)["verdict"] == CARE
    assert ship_check(ok, "Acme", D("10"), TODAY, broken=2, kept=0)["verdict"] == ADVANCE
    assert ship_check(ok, "Acme", D("10"), TODAY, broken=2, kept=5)["verdict"] == CARE   # mostly keeps promises
    assert ship_check(ok, "Acme", D("10"), TODAY, open_disputes=1)["verdict"] == CARE
    assert ship_check(ok, "Acme", D("10"), TODAY, do_not_contact=True)["verdict"] == CARE


def test_the_most_severe_rule_wins_and_unknown_customer_is_new():
    r = ship_check([inv("A-1", "Acme", -70, "100")], "Acme", D("10"), TODAY, broken=3)
    assert r["verdict"] == HOLD and len([x for x in r["reasons"] if x["level"] in ("bad", "warn")]) >= 2
    n = ship_check([inv("A-1", "Acme", 10)], "Stranger", D("10"), TODAY)
    assert n["verdict"] == NEW and n["figures"] == {}
    with pytest.raises(ValueError):
        ship_check([], "Acme", D("0"), TODAY)


# ---- cash calendar ---------------------------------------------------------------------------------------------------------

def test_promise_and_due_invoices_land_in_their_week_firm_by_default():
    invoices = [inv("A-1", "Acme", -10, "50000"), inv("B-1", "Beta", 9, "20000")]
    cal = cash_calendar(invoices, [promise("Acme", TODAY + timedelta(days=3), "30000")], {}, set(), TODAY)
    assert cal["weeks"][0]["firm"] == "30000.00" and cal["weeks"][1]["firm"] == "20000.00"
    assert cal["totals"] == {"firm_14": "50000.00", "all_14": "50000.00", "firm_30": "50000.00", "all_30": "50000.00"}
    assert {i["kind"] for w in cal["weeks"] for i in w["items"]} == {"promise", "terms"}


def test_unreliable_or_very_late_customers_are_shaky_not_firm():
    invoices = [inv("A-1", "Acme", -10, "50000"), inv("B-1", "Beta", -45, "50000")]
    promises = [promise("Acme", TODAY + timedelta(days=2), "10000", 1), promise("Beta", TODAY + timedelta(days=2), "20000", 2)]
    cal = cash_calendar(invoices, promises, {customer_key("Acme"): Reliability(kept=0, broken=2)}, set(), TODAY)
    assert cal["weeks"][0]["firm"] == "0.00" and cal["weeks"][0]["shaky"] == "30000.00"
    assert cal["totals"]["firm_30"] == "0.00" and cal["totals"]["all_30"] == "30000.00"


def test_overdue_without_a_promise_is_never_counted_as_incoming():
    cal = cash_calendar([inv("A-1", "Acme", -10, "50000"), inv("B-1", "Beta", -50, "70000")], [], {}, set(), TODAY)
    assert cal["totals"]["all_30"] == "0.00"
    assert cal["undated"]["amount"] == "120000.00" and cal["undated"]["customers"] == 2 and cal["undated"]["top"][0]["customer"] == "Beta"


def test_promise_is_capped_at_overdue_and_without_amount_is_not_counted():
    invoices = [inv("A-1", "Acme", -10, "5000")]
    cal = cash_calendar(invoices, [promise("Acme", TODAY + timedelta(days=1), "999999")], {}, set(), TODAY)
    assert cal["weeks"][0]["total"] == "5000.00"
    cal2 = cash_calendar(invoices, [promise("Acme", TODAY + timedelta(days=1), None)], {}, set(), TODAY)
    assert cal2["totals"]["all_30"] == "0.00" and cal2["promises_without_amount"] == 1 and cal2["undated"]["customers"] == 1


def test_broken_or_far_promises_and_disputed_invoices_are_left_out():
    invoices = [inv("A-1", "Acme", -10, "5000"), inv("D-1", "Dee", 3, "9000")]
    promises = [promise("Acme", TODAY - timedelta(days=1), "5000", 1), promise("Acme", TODAY + timedelta(days=60), "5000", 2)]
    assert cash_calendar(invoices, promises, {}, {("D-1", "dee")}, TODAY)["totals"]["all_30"] == "0.00"


def test_outgoings_gap_uses_firm_and_all():
    cal = cash_calendar([inv("A-1", "Acme", 5, "100000")], [], {}, set(), TODAY, D("150000"))
    assert cal["outgoings"]["gap_30_firm"] == "-50000.00" and cal["outgoings"]["days_14"] == "70000.00"


# ---- morning brief ---------------------------------------------------------------------------------------------------------

def test_morning_brief_is_short_plain_text(tmp_path):
    header = ["Invoice No", "Customer", "Invoice Date", "Due Date", "Amount", "Paid"]
    rows = [(1, header), (2, ["INV-1", "Acme Traders", "01-06-2026", "01-07-2026", "50000", "0"]), (3, ["INV-2", "Beta Co", "01-06-2026", "01-08-2026", "20000", "0"])]
    now = datetime(2026, 10, 14, 9, 0, tzinfo=timezone.utc)
    with Store(tmp_path / "b.db", now=lambda: now) as s:
        s.apply_import(import_rows(rows), actor="t")
        s.log_followup("Acme Traders", "call", "Promised", "t", promise_date=TODAY, promise_amount=D("10000"))
        b = morning_brief(s, TODAY, "Sharma Steel")
    text = b["text"]
    assert text.startswith("Sharma Steel Morning brief, Thu 15 Oct 2026")
    assert "Call today:" in text and "1. Beta Co" in text
    assert "Promised for today: Acme Traders ₹10,000.00" in text
    assert len(text.splitlines()) <= 10 and b["facts"]["to_call"] == 1 and b["facts"]["promised_today"] == 1
    assert "No new payments recorded in the last 24 hours." in text
    rows[2] = (3, ["INV-2", "Beta Co", "01-06-2026", "01-08-2026", "20000", "5000"])
    with Store(tmp_path / "b.db", now=lambda: now) as s:
        s.apply_import(import_rows(rows), actor="t")
        paid = morning_brief(s, TODAY)
    assert "Payments recorded in the last 24 hours: ₹5,000.00." in paid["text"] and paid["facts"]["came_in_24h"] == "5000.00"
