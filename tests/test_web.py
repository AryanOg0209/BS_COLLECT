import http.client
import json
import re
import threading
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from collectos import web
from collectos.cli import main
from collectos.importer import import_rows
from collectos.store import Store

HEADER = ["Invoice No", "Customer", "Invoice Date", "Due Date", "Amount", "Paid"]
AS_OF = date(2026, 10, 15)
NOW = datetime(2026, 10, 15, 9, 0, tzinfo=timezone.utc)


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "t.db"
    rows = [
        (1, HEADER),
        (2, ["A1", "Acme Traders", "01-06-2026", "01-07-2026", "100000", "0"]),   # 106 days overdue
        (3, ["A2", "Acme Traders", "01-09-2026", "01-11-2026", "5000", "0"]),     # not due
        (4, ["B1", "Beta Co", "01-09-2026", "01-10-2026", "20000", "5000"]),      # 14 days overdue, part paid
        (5, ["C1", "<b>Evil</b> & Co", "01-09-2026", "01-10-2026", "300", "0"]),  # markup in a name stays data
    ]
    with Store(path, now=lambda: NOW) as s:
        s.apply_import(import_rows(rows), actor="t")
        s.take_snapshot(AS_OF, actor="t")
        s.log_followup("Beta Co", "call", "Will pay Friday", actor="t", promise_date=date(2026, 10, 20))
    return path


@pytest.fixture
def server(db):
    srv = web.make_server(db)
    thread = threading.Thread(target=lambda: srv.serve_forever(poll_interval=0.02), daemon=True)
    thread.start()
    yield srv
    srv.shutdown()
    srv.server_close()


def request(srv, path, method="GET", host=None, body=None):
    conn = http.client.HTTPConnection("127.0.0.1", srv.server_address[1], timeout=10)
    conn.putrequest(method, path, skip_host=True)
    conn.putheader("Host", f"127.0.0.1:{srv.server_address[1]}" if host is None else host)
    conn.putheader("Content-Length", str(len(body or b"")))
    conn.endheaders(body)
    resp = conn.getresponse()
    data = resp.read()
    headers = {k.lower(): v for k, v in resp.getheaders()}
    conn.close()
    return resp.status, headers, data


def get_json(srv, path, **kw):
    status, headers, data = request(srv, path, **kw)
    return status, json.loads(data)


# ---- data -----------------------------------------------------------------------------


def test_overview_numbers_match_the_ledger(server):
    status, d = get_json(server, "/api/overview?as_of=2026-10-15")
    assert status == 200 and d["as_of"] == "2026-10-15" and d["db_name"] == "t.db"
    k = d["kpi"]
    assert k["total_outstanding"] == "120300.00" and k["overdue"] == "115300.00"   # 100000 + 15000 + 300 ; not 5000 (not due)
    assert (k["open_invoices"], k["overdue_invoices"], k["customers_overdue"]) == (4, 3, 3)
    assert k["overdue_share"] == f"{Decimal('115300') / Decimal('120300'):.4f}"
    assert k["waiting"] == 1 and k["to_contact_today"] == 2   # Beta promised a date, so it waits


def test_money_is_sent_as_exact_strings_never_floats(server):
    _, d = get_json(server, "/api/overview?as_of=2026-10-15")
    def walk(x):
        if isinstance(x, dict):
            for v in x.values():
                yield from walk(v)
        elif isinstance(x, list):
            for v in x:
                yield from walk(v)
        else:
            yield x
    assert not [v for v in walk(d) if isinstance(v, float)]
    assert all(re.fullmatch(r"-?\d+\.\d{2}", a["amount"]) for a in d["ageing"])


def test_ageing_queue_promises_customers_and_trend(server):
    _, d = get_json(server, "/api/overview?as_of=2026-10-15")
    amounts = {a["bucket"]: (a["amount"], a["count"]) for a in d["ageing"]}
    assert amounts["Current"] == ("5000.00", 1) and amounts["91+"] == ("100000.00", 1) and amounts["1-30"] == ("15300.00", 2)
    assert [q["customer"] for q in d["queue"]] == ["Acme Traders", "<b>Evil</b> & Co"]
    assert d["queue"][0]["action"] == "Call" and any("overdue" in r for r in d["queue"][0]["reasons"])
    assert d["waiting"][0]["customer"] == "Beta Co" and "promised to pay by 20-Oct-2026" in d["waiting"][0]["reason"]
    assert d["promises"][0]["state"] == "pending"
    assert d["customers"][0]["customer"] == "Acme Traders" and d["customers"][0]["overdue"] == "100000.00"
    assert len(d["trend"]) == 1 and d["trend"][0]["as_of"] == "2026-10-15"
    assert d["followups"][0]["outcome"] == "Will pay Friday"


def test_promise_turns_broken_when_the_date_has_passed(server):
    _, d = get_json(server, "/api/overview?as_of=2026-10-21")
    assert d["promises"][0]["state"] == "broken" and d["kpi"]["broken_promises"] == 1


def test_markup_in_names_is_returned_as_plain_data(server):
    _, d = get_json(server, "/api/overview?as_of=2026-10-15")
    assert "<b>Evil</b> & Co" in [c["customer"] for c in d["customers"]]  # escaped at render time by textContent


def test_bad_date_is_a_400_not_a_crash(server):
    status, d = get_json(server, "/api/overview?as_of=not-a-date")
    assert status == 400 and "error" in d


# ---- drafts and report ----------------------------------------------------------------


def test_draft_endpoint(server):
    status, d = get_json(server, "/api/draft?customer=acme%20traders&sender=Rohit&company=Sharma&as_of=2026-10-15")
    assert status == 200 and d["template"] == "firm" and "Rohit" in d["email"] and "INV" not in d["email"] and "A1" in d["email"]
    assert get_json(server, "/api/draft?customer=Nobody")[0] == 404
    assert get_json(server, "/api/draft")[0] == 400
    assert get_json(server, "/api/draft?customer=Beta%20Co&as_of=2026-10-15")[0] == 200


def test_report_download_is_a_valid_workbook(server, tmp_path):
    status, headers, data = request(server, "/report.xlsx?company=Test%20Co&as_of=2026-10-15")
    assert status == 200 and "attachment" in headers["content-disposition"] and "2026-10-15" in headers["content-disposition"]
    out = tmp_path / "r.xlsx"
    out.write_bytes(data)
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames and "Trend" in wb.sheetnames
    assert wb["Summary"]["A1"].value == "Receivables summary: Test Co"


# ---- safety properties ----------------------------------------------------------------


def test_binds_to_loopback_only(server):
    assert server.server_address[0] == "127.0.0.1"


@pytest.mark.parametrize("host", ["evil.example", "evil.example:80", "192.168.1.5:8765", "", "127.0.0.1:1"])
def test_foreign_host_headers_are_rejected(server, host):
    status, _, _ = request(server, "/api/overview", host=host)
    assert status == 403


def test_localhost_name_is_accepted(server):
    status, _, _ = request(server, "/api/overview?as_of=2026-10-15", host=f"localhost:{server.server_address[1]}")
    assert status == 200


@pytest.mark.parametrize("method", ["POST", "PUT", "DELETE", "PATCH"])
def test_server_is_read_only(server, db, method):
    before = Path(db).read_bytes()
    status, headers, _ = request(server, "/api/overview", method=method, body=b"{}")
    assert status == 405 and "GET" in headers["allow"]
    with Store(db) as s:  # nothing changed
        assert len(s.invoices()) == 4 and len(s.followups()) == 1
    assert before  # database file still readable


@pytest.mark.parametrize("path", ["/static/../web.py", "/static/..%2fweb.py", "/static/%2e%2e/store.py", "/static/app.js/../../store.py",
                                  "/static/nope.js", "/static/", "/etc/passwd", "/../../etc/passwd", "/static/app.css%00.py"])
def test_no_path_traversal_or_unlisted_files(server, path):
    status, _, data = request(server, path)
    assert status == 404 and b"import" not in data


def test_security_headers_on_every_response(server):
    for path in ("/", "/static/app.js", "/api/overview?as_of=2026-10-15", "/nope"):
        _, h, _ = request(server, path)
        assert "default-src 'none'" in h["content-security-policy"] and "script-src 'self'" in h["content-security-policy"]
        assert h["x-content-type-options"] == "nosniff" and h["x-frame-options"] == "DENY"
        assert h["cache-control"] == "no-store" and h["referrer-policy"] == "no-referrer"
        assert "access-control-allow-origin" not in h  # no CORS: other sites cannot read the API


def test_static_assets_are_served(server):
    status, h, data = request(server, "/")
    assert status == 200 and h["content-type"].startswith("text/html") and b"CollectOS" in data
    assert request(server, "/static/app.js")[1]["content-type"].startswith("text/javascript")
    assert request(server, "/static/app.css")[1]["content-type"].startswith("text/css")
    assert request(server, "/static/logo.svg")[1]["content-type"] == "image/svg+xml"


def test_head_requests_have_no_body(server):
    status, h, data = request(server, "/", method="HEAD")
    assert status == 200 and data == b"" and int(h["content-length"]) > 0


def test_missing_database_is_refused_not_created(tmp_path):
    with pytest.raises(FileNotFoundError, match="database not found"):
        web.make_server(tmp_path / "missing.db")
    assert not (tmp_path / "missing.db").exists()


def test_ui_command_reports_a_missing_database(capsys, tmp_path):
    assert main(["ui", "--db", str(tmp_path / "nope.db"), "--no-open"]) == 1
    assert "database not found" in capsys.readouterr().err


# ---- guard rails on the front-end source -------------------------------------------------

WEBUI = Path(web.__file__).parent / "webui"


def test_frontend_never_builds_html_from_strings():
    js = (WEBUI / "app.js").read_text() + (WEBUI / "features.js").read_text()
    for banned in ("innerHTML", "outerHTML", "insertAdjacentHTML", "document.write", "eval(", "new Function", "setAttribute(\"style\""):
        assert banned not in js, f"{banned} would let imported text become markup or code"


# Links a person clicks that open their own app. They navigate away; the page never loads or fetches them.
ALLOWED_CLICK_TARGETS = {"http://www.w3.org/2000/svg", "https://wa.me/"}


def test_frontend_loads_nothing_from_other_origins():
    for name in ("index.html", "app.css", "app.js", "features.js", "logo.svg"):
        text = (WEBUI / name).read_text()
        urls = [u for u in re.findall(r"https?://[^\s\"')`$]+", text) if u not in ALLOWED_CLICK_TARGETS]
        assert not urls, f"{name} references external resources: {urls}"
    assert "style=" not in (WEBUI / "index.html").read_text()  # inline styles are blocked by the CSP
    assert "@import" not in (WEBUI / "app.css").read_text()


def test_frontend_only_ever_fetches_its_own_server():
    for name in ("app.js", "features.js"):
        text = (WEBUI / name).read_text()
        assert not re.search(r"fetch\(\s*[\"'`]https?:", text), f"{name} fetches an absolute URL"
        assert not re.search(r"(XMLHttpRequest|WebSocket|EventSource|sendBeacon|importScripts)", text), f"{name} opens another kind of connection"
        for src in re.findall(r"\bsrc\s*[:=]\s*[\"']([^\"']+)", text):
            assert src.startswith("/static/") or src.startswith("data:"), f"{name} loads {src}"
