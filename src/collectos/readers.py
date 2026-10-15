"""File readers: CSV and XLSX to raw rows. No interpretation here."""

import csv
from pathlib import Path

from openpyxl import load_workbook

# Each row is (1-based row number in the source file, list of cell values).
RawRows = list[tuple[int, list[object]]]


class UnsupportedFile(ValueError):
    pass


def read_rows(path: str | Path) -> RawRows:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise UnsupportedFile(f"unsupported file type {suffix!r}; use .csv or .xlsx (save .xls as .xlsx)")


def _read_csv(path: Path) -> RawRows:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [(i, list(row)) for i, row in enumerate(csv.reader(handle), start=1)]


def _read_xlsx(path: Path) -> RawRows:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [(i, list(row)) for i, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    finally:
        workbook.close()
