"""Owner report: one Excel workbook summarising receivables and today's actions.

Values are written as computed numbers (not formulas) so the file never depends on
recalculation. Text cells are forced to string type so imported names such as
"=HYPERLINK(...)" can never execute as spreadsheet formulas.
"""

from collections import Counter
from collections.abc import Iterable, Sequence
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .ageing import CURRENT, DEFAULT_EDGES, NO_DUE_DATE, Ageing, Status, assess, bucket_labels, summarise_by_customer
from .models import Invoice, Promise
from .workqueue import Queue

if TYPE_CHECKING:
    from .store import FollowUp

INR_FORMAT = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00"
DATE_FORMAT = "DD-MMM-YYYY"
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ZERO = Decimal("0.00")

DISCLAIMER = (
    "This report describes the data supplied on the date shown. It does not predict or guarantee "
    "collections, and it is not accounting, tax or legal advice."
)


def write_owner_report(
    path: str | Path,
    *,
    company: str,
    as_of: date,
    invoices: Iterable[Invoice],
    queue: Queue,
    open_promises: Sequence[Promise] = (),
    followups: Sequence["FollowUp"] = (),
    warnings: Sequence[str] = (),
    edges: tuple[int, ...] = DEFAULT_EDGES,
) -> Path:
    results = [assess(i, as_of, edges) for i in invoices]
    wb = Workbook()
    _summary(wb.active, company, as_of, results, queue, open_promises, followups, edges)
    _queue_sheet(wb.create_sheet("Today's Queue"), queue)
    _waiting_sheet(wb.create_sheet("Waiting"), queue)
    _customers_sheet(wb.create_sheet("Ageing by Customer"), results, edges)
    _invoices_sheet(wb.create_sheet("Open Invoices"), results)
    _promises_sheet(wb.create_sheet("Promises"), open_promises, as_of)
    _notes_sheet(wb.create_sheet("Notes"), warnings, as_of, edges)
    path = Path(path)
    wb.save(path)
    return path


# ---- sheets ---------------------------------------------------------------------


def _summary(ws, company, as_of, results, queue, promises, followups, edges) -> None:
    ws.title = "Summary"
    open_ = [r for r in results if r.status not in (Status.PAID, Status.OVERPAID)]
    outstanding = sum((r.outstanding for r in open_), _ZERO)
    overdue = sum((r.outstanding for r in open_ if r.status is Status.OVERDUE), _ZERO)
    window_start = as_of - timedelta(days=6)
    recent = sum(1 for f in followups if window_start <= datetime.fromisoformat(f.ts).date() <= as_of)
    broken = sum(1 for p in promises if p.promised_date < as_of)
    overpaid = sum(1 for r in results if r.status is Status.OVERPAID)

    _put(ws, 1, 1, f"Receivables summary: {company}", Font(bold=True, size=14))
    _put(ws, 2, 1, f"As of {as_of:%d-%b-%Y}")
    rows = [
        ("Total outstanding", outstanding, True),
        ("Overdue", overdue, True),
        ("Overdue share of outstanding", None if outstanding == 0 else f"{overdue / outstanding:.0%}", False),
        ("Open invoices", len(open_), False),
        ("Customers with overdue balance", len({r.invoice.customer for r in open_ if r.status is Status.OVERDUE}), False),
        ("Customers to contact today", len(queue.work), False),
        ("Customers waiting (promise pending / recently contacted)", len(queue.waiting), False),
        ("Open promises past their date (broken)", broken, False),
        ("Follow-ups logged in the last 7 days", recent, False),
        ("Invoices overpaid (needs review)", overpaid, False),
    ]
    for i, (label, value, money) in enumerate(rows, start=4):
        _put(ws, i, 1, label)
        _put(ws, i, 2, value, money=money)

    start = 4 + len(rows) + 1
    _put(ws, start, 1, "Ageing", Font(bold=True))
    buckets = [CURRENT, *bucket_labels(edges), NO_DUE_DATE]
    count, amount = Counter(), {b: _ZERO for b in buckets}
    for r in open_:
        count[r.bucket] += 1
        amount[r.bucket] += r.outstanding
    _table(ws, start + 1, ["Bucket (days overdue)", "Invoices", "Amount"], [[b, count[b], amount[b]] for b in buckets], {3})

    top_row = start + 2 + len(buckets) + 1
    _put(ws, top_row, 1, "Top 10 customers by overdue amount", Font(bold=True))
    top = sorted(summarise_by_customer(results).values(), key=lambda s: (-s.overdue, s.customer.casefold()))
    top = [s for s in top if s.overdue > 0][:10]
    _table(ws, top_row + 1, ["Customer", "Overdue", "Oldest (days)"], [[s.customer, s.overdue, s.max_days_overdue] for s in top], {2})
    _widths(ws, [52, 18, 16])


def _queue_sheet(ws, queue: Queue) -> None:
    rows = [
        [n, i.customer, i.action, i.score, i.overdue, i.outstanding, i.max_days_overdue,
         "never" if i.days_since_contact is None else i.days_since_contact, "; ".join(i.reasons)]
        for n, i in enumerate(queue.work, start=1)
    ]
    _table(ws, 1, ["#", "Customer", "Action", "Score", "Overdue", "Outstanding", "Oldest (days)", "Days since contact", "Why"],
           rows, {5, 6})
    _widths(ws, [5, 34, 15, 8, 16, 16, 14, 18, 90])


def _waiting_sheet(ws, queue: Queue) -> None:
    rows = [[w.customer, w.overdue, w.reason] for w in queue.waiting]
    rows += [[w.customer, None, w.reason] for w in queue.promises_to_review]
    _table(ws, 1, ["Customer", "Overdue", "Why not in today's queue"], rows, {2})
    _widths(ws, [34, 16, 80])


def _customers_sheet(ws, results: list[Ageing], edges) -> None:
    buckets = [CURRENT, *bucket_labels(edges), NO_DUE_DATE]
    summaries = sorted(summarise_by_customer(results).values(), key=lambda s: (-s.overdue, s.customer.casefold()))
    rows = [
        [s.customer, s.outstanding, s.overdue, *[s.by_bucket.get(b, _ZERO) for b in buckets], s.open_invoices, s.max_days_overdue]
        for s in summaries
    ]
    money = set(range(2, 4 + len(buckets)))
    _table(ws, 1, ["Customer", "Outstanding", "Overdue", *buckets, "Open invoices", "Oldest (days)"], rows, money)
    _widths(ws, [34] + [15] * (2 + len(buckets)) + [14, 14])


def _invoices_sheet(ws, results: list[Ageing]) -> None:
    shown = sorted(
        (r for r in results if r.status is not Status.PAID),
        key=lambda r: (-r.days_overdue, r.invoice.customer.casefold(), r.invoice.invoice_no),
    )
    rows = [
        [r.invoice.invoice_no, r.invoice.customer, r.invoice.invoice_date, r.invoice.due_date, r.invoice.amount,
         r.invoice.paid, r.outstanding, r.status.value.replace("_", " "), r.days_overdue, r.bucket, "yes" if r.is_partial else ""]
        for r in shown
    ]
    _table(ws, 1, ["Invoice", "Customer", "Invoice date", "Due date", "Amount", "Paid", "Outstanding", "Status",
                   "Days overdue", "Bucket", "Part-paid"], rows, {5, 6, 7})
    _widths(ws, [16, 34, 14, 14, 15, 15, 15, 14, 13, 13, 10])


def _promises_sheet(ws, promises: Sequence[Promise], as_of: date) -> None:
    rows = [
        [p.customer, p.promised_date, p.promised_amount, "BROKEN (date passed)" if p.promised_date < as_of else "pending"]
        for p in sorted(promises, key=lambda p: (p.promised_date, p.customer.casefold()))
    ]
    _table(ws, 1, ["Customer", "Promised date", "Promised amount", "State"], rows, {3})
    _widths(ws, [34, 16, 18, 24])


def _notes_sheet(ws, warnings: Sequence[str], as_of: date, edges) -> None:
    lines = [
        "How to read this report",
        f"- Ageing is measured from each invoice's due date to {as_of:%d-%b-%Y}. Buckets: Current (not yet due), "
        + ", ".join(bucket_labels(edges)) + " days overdue.",
        f"- '{NO_DUE_DATE}' invoices cannot be aged and are excluded from Overdue until a due date is supplied.",
        "- Queue scores are the sum of the visible reasons; nothing is hidden or predicted.",
        "- A customer who has an unexpired promise, or was contacted very recently, is shown under Waiting.",
        "- A promise is broken the day after its promised date if it was not marked kept.",
        "",
        DISCLAIMER,
    ]
    if warnings:
        lines += ["", "Data warnings from the import"] + [f"- {w}" for w in warnings]
    headings = {lines[0], "Data warnings from the import"}
    for i, line in enumerate(lines, start=1):
        _put(ws, i, 1, line, Font(bold=True) if line in headings else None)
    _widths(ws, [140])


# ---- helpers --------------------------------------------------------------------


def _table(ws, top: int, headers: list[str], rows: list[list[object]], money_cols: set[int]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = _put(ws, top, c, h, _HEADER_FONT)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start=top + 1):
        for c, value in enumerate(row, start=1):
            _put(ws, r, c, value, money=c in money_cols)
    if top == 1:
        ws.freeze_panes = "A2"


def _put(ws, row: int, col: int, value: object, font: Font | None = None, money: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # never let imported text act as a formula
    if font:
        cell.font = font
    if isinstance(value, Decimal) or money:
        cell.number_format = INR_FORMAT
    elif isinstance(value, date):
        cell.number_format = DATE_FORMAT
    return cell


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
